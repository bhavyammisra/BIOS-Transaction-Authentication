"""Visual (v3) Pramaan deck: less text, bigger type, icons and diagrams.
Usage: python build_deck_v3.py <v2 deck.pptx> <out.pptx> <icons dir> [--n8n]   (--n8n builds v4)
Starts from the v2 deck (keeps the BIOS template shapes and speaker notes) and redraws the content.
Icons are PNGs named <icon>_<HEX>.png (rendered from react-icons)."""
import sys, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from build_deck import (NAVY, ICE, PANEL, TEXT, MUTED, GREEN, GREEN_T, AMBER, AMBER_T, RED, RED_T, WHITE,
                        BODY, HEAD, rgb, remove, add_par, textbox, box, pill, arrow, label, heading, sources)

ICON_DIR = None
N8N = False  # v4: show where n8n is used (--n8n)
N8N_C, N8N_T = "EA4B71", "FDE8EE"  # n8n brand coral + tint: every n8n mention in v4 is highlighted
L, C = PP_ALIGN.LEFT, PP_ALIGN.CENTER

def icon(slide, name, x, y, size, color=NAVY):
    return slide.shapes.add_picture(os.path.join(ICON_DIR, f"{name}_{color}.png"), Inches(x), Inches(y), Inches(size), Inches(size))

def badge(slide, name, x, y, d, fill=NAVY, color=WHITE, ratio=0.56):
    """Icon centred in a filled circle: the deck's visual motif."""
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    s.name = f"Badge {name}"
    i = d * ratio
    icon(slide, name, x + (d - i) / 2, y + (d - i) / 2, i, color)

def bar(slide, x, y, w, h, fill, name):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill); s.line.fill.background(); s.shadow.inherit = False
    s.name = name
    return s

def text(slide, x, y, w, h, name, runs, align=L, anchor=MSO_ANCHOR.TOP):
    tb, tf = textbox(slide, x, y, w, h, name, anchor=anchor)
    add_par(tf, runs, align=align, first=True)
    return tf

# ---------------------------------------------------------------- slide 1
def slide1(s):
    heading(s, headline="> Fraud has moved past OTPs: victims now approve the payment themselves")
    st, tf = box(s, 4.15, 0.66, 8.7, 0.9, "Team strip", fill=PANEL, radius=0.12, pad=0.12)
    add_par(tf, [("Team Caffeinated Nerds", 11.5, True, NAVY), ("   ·   Challenge: ", 11, False, NAVY), ("Frontier · PS-04", 11, True, NAVY),
                 (" · Fast and Secure Online Transaction Authentication", 11, False, NAVY)], first=True, align=L)
    for i, (name, roll) in enumerate([("Bhavyam Misra", "1024030535"), ("Gunagye Jain", "1024030019"),
                                      ("Abhijay Nagal", "1024170013"), ("Shefali Sharma", "1024030284")]):
        tb, ttf = textbox(s, 4.27 + i * 2.13, 1.01, 2.05, 0.45, f"Member {i+1}")
        add_par(ttf, [(name, 10.5, True, TEXT)], first=True, align=L)
        add_par(ttf, [("Roll No: " + roll, 10, False, MUTED)], align=L)

    # Asha's story as a 4-step strip
    text(s, 0.5, 2.8, 6.75, 0.3, "Story label", [("THE PROBLEM: ASHA, 72, PAYS A SCAMMER WITH HER OWN PIN", 11, True, NAVY)])
    steps = [("phonecall", AMBER, "Fake “CBI” officer calls"), ("phone", NAVY, "Opens UPI on her own phone"),
             ("lock", NAVY, "Types her own PIN"), ("checkc", RED, "Approved: ₹85,000 gone")]
    for i, (ic, col, cap) in enumerate(steps):
        x = 0.5 + i * 1.75
        b, _ = box(s, x, 3.12, 1.5, 1.35, f"Step {i+1}", fill=RED_T if i == 3 else PANEL, radius=0.1)
        badge(s, ic, x + 0.45, 3.22, 0.6, fill=col)
        text(s, x + 0.08, 3.86, 1.34, 0.55, f"Step {i+1} text", [(cap, 11.5, True, RED if i == 3 else TEXT)], align=C)
        if i < 3:
            arrow(s, x + 1.52, 3.8, x + 1.73, 3.8, f"Step arrow {i+1}")
    # the core insight
    b, tf = box(s, 0.5, 4.62, 6.75, 0.72, "Insight", fill=NAVY, radius=0.14, pad=0.2, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tf, [("A PIN checks ", 17, False, WHITE, False, HEAD), ("who", 17, True, ICE, False, HEAD), (" you are. Nobody checks ", 17, False, WHITE, False, HEAD),
                 ("what's happening", 17, True, ICE, False, HEAD), (".", 17, False, WHITE, False, HEAD)], first=True, align=C)
    # who we serve
    text(s, 0.5, 5.46, 6.75, 0.3, "Stakeholder label", [("WHO WE SERVE", 11, True, NAVY)])
    who = [("bank", "Banks, UPI apps & gateways"), ("users", "Seniors & first-time payers"), ("clock", "Decide inside UPI's 15-s window")]
    for i, (ic, cap) in enumerate(who):
        x = 0.5 + i * 2.3
        b, _ = box(s, x, 5.76, 2.15, 0.58, f"Who {i+1}", fill=PANEL, radius=0.2)
        badge(s, ic, x + 0.08, 5.83, 0.44)
        text(s, x + 0.6, 5.76, 1.5, 0.58, f"Who {i+1} text", [(cap, 11, True, TEXT)], anchor=MSO_ANCHOR.MIDDLE)

    # evidence with icons and big numbers
    text(s, 7.55, 2.8, 5.3, 0.3, "Evidence label", [("EVIDENCE", 11, True, NAVY)])
    stats = [("rupee", "₹22,495 cr", "lost to cyber fraud in India in 2025 (MHA)"),
             ("phonecall", "~85%", "of losses: scams where the victim pays (MHA)"),
             ("trend", "98.5%", "of fraud value is in payments above ₹10k (RBI)")]
    for i, (ic, big, small) in enumerate(stats):
        y = 3.12 + i * 1.08
        b, _ = box(s, 7.55, y, 5.3, 0.96, f"Stat {i+1}", fill=PANEL, radius=0.12)
        badge(s, ic, 7.72, y + 0.18, 0.6)
        text(s, 8.5, y + 0.06, 4.2, 0.5, f"Stat {i+1} value", [(big, 26, True, NAVY, False, HEAD)])
        text(s, 8.5, y + 0.55, 4.25, 0.35, f"Stat {i+1} text", [(small, 12, False, TEXT)])

    sdgs = [("00689D", "PRIMARY · SDG 16", "16.4: cut illicit financial flows & organised crime"),
            ("A21942", "SUPPORTING · SDG 8", "8.10: access to financial services for all"),
            ("DD5A1F", "SUPPORTING · SDG 9", "9.1: reliable, resilient (digital) infrastructure")]
    for i, (col, head, body) in enumerate(sdgs):
        b, tf = box(s, 0.5 + i * 4.175, 6.48, 4.0, 0.6, f"SDG {i+1}", fill=col, radius=0.18, pad=0.12, anchor=MSO_ANCHOR.MIDDLE)
        add_par(tf, [(head, 10.5, True, WHITE)], first=True, align=L)
        add_par(tf, [(body, 10.5, False, WHITE)], align=L)
    sources(s, "Sources: MHA/NCRP data via Parliament replies (2024, 2025); RBI Discussion Paper on safeguards in digital payments, 9 Apr 2026. Links: repo › docs/02-research/problem-evidence.md")

# ---------------------------------------------------------------- slide 2
def slide2(s):
    heading(s, headline="> Today's defences check the credential or the receiver, never the moment")
    cols = [("lock", "Static 2FA", "UPI PIN · OTP · 3-D Secure", PANEL),
            ("cpu", "Risk engines", "Stripe Radar · Sift · BioCatch", PANEL),
            ("bank", "India's tools", "NPCI FRM · DoT FRI · MuleHunter", PANEL),
            ("shield", "Pramaan", "this project", NAVY)]
    x0, lw, cw, gap = 0.5, 3.55, 2.13, 0.07
    rows = [("Checks who holds the credential", "yynY"),
            ("Scores the receiver's reputation", "nyyY"),
            ("Reads the live moment: call, device, typing", "npnY"),
            ("Explains the decision to the user", "nnnY"),
            ("Picks the right challenge for the risk", "nnnY"),
            ("Affordable for small banks & apps", "ynyY")]
    # header row
    for j, (ic, name, ex, fill) in enumerate(cols):
        x = x0 + lw + j * (cw + gap)
        dark = fill == NAVY
        b, _ = box(s, x, 2.8, cw, 1.08, f"Col {j+1}", fill=fill, radius=0.1)
        badge(s, ic, x + (cw - 0.46) / 2, 2.87, 0.46, fill=WHITE if dark else NAVY, color=NAVY if dark else WHITE)
        text(s, x, 3.36, cw, 0.28, f"Col {j+1} name", [(name, 13, True, WHITE if dark else NAVY, False, HEAD)], align=C)
        text(s, x, 3.62, cw, 0.24, f"Col {j+1} ex", [(ex, 9, False, ICE if dark else MUTED, True)], align=C)
    # ticks matrix
    mark = {"y": ("check", GREEN), "Y": ("check", GREEN), "n": ("x", RED), "p": ("minus", AMBER)}
    for i, (lab, vals) in enumerate(rows):
        y = 3.98 + i * 0.42
        fill = PANEL if i % 2 == 0 else WHITE
        bar(s, x0, y, lw + 3 * (cw + gap) + cw, 0.38, fill, f"Row {i+1} bg")
        text(s, x0 + 0.12, y, lw - 0.15, 0.38, f"Row {i+1} label", [(lab, 12.5, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
        for j, v in enumerate(vals):
            x = x0 + lw + j * (cw + gap)
            if j == 3:
                bar(s, x, y, cw, 0.38, "DDE3F7", f"Row {i+1} hl")
            ic, col = mark[v]
            icon(s, ic, x + (cw - 0.28) / 2, y + 0.05, 0.28, col)
    # legend
    lx = 0.62
    for ic, col, t in (("check", GREEN, "yes"), ("minus", AMBER, "partly"), ("x", RED, "no")):
        icon(s, ic, lx, 6.55, 0.2, col)
        text(s, lx + 0.24, 6.52, 0.7, 0.26, f"Legend {t}", [(t, 10, False, MUTED)])
        lx += 0.95
    b, tf = box(s, 3.55, 6.5, 9.3, 0.6, "Gap statement", fill=NAVY, radius=0.2, pad=0.18, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tf, [("The gap: ", 13.5, True, ICE, False, HEAD),
                 ("nothing lightweight reads the moment and answers with an explained, right-sized challenge.", 13, False, WHITE)], first=True, align=L)
    sources(s, "Sources: RBI (25 Sep 2025, 9 Apr 2026); Stripe Radar, Sift, BioCatch docs; NPCI FRM; DoT FRI; MuleHunter.AI. Links: repo › docs/02-research/prior-art.md", y=7.14)

# ---------------------------------------------------------------- slide 3
def slide3(s):
    heading(s, headline="> Pramaan: a PIN checks who you are. We check what's happening, before every payment.")
    text(s, 0.5, 2.8, 5.75, 0.3, "Build label", [("WHAT WE BUILD: A RISK CHECK BEFORE THE PIN", 11, True, NAVY)])
    flow = [("phone", NAVY, "Bank / UPI app calls us"), ("shield", NAVY, "Scores 20+ live signals"), ("zap", GREEN, "Decision in ~40 ms")]
    for i, (ic, col, cap) in enumerate(flow):
        x = 0.5 + i * 2.05
        b, _ = box(s, x, 3.12, 1.65, 1.2, f"Flow {i+1}", fill=PANEL, radius=0.1)
        badge(s, ic, x + 0.53, 3.2, 0.6, fill=col)
        text(s, x + 0.06, 3.84, 1.53, 0.45, f"Flow {i+1} text", [(cap, 11.5, True, TEXT)], align=C)
        if i < 2:
            arrow(s, x + 1.67, 3.72, x + 2.03, 3.72, f"Flow arrow {i+1}")
    text(s, 0.5, 4.48, 5.75, 0.3, "Diff label", [("WHY IT'S DIFFERENT", 11, True, NAVY)])
    diffs = [("user", "Judged against each user's own habits"),
             ("rupee", "Amount-aware: a ₹40 tea is never stopped"),
             ("msg", "A plain-language reason with every decision")]
    for i, (ic, t) in enumerate(diffs):
        y = 4.8 + i * 0.52
        badge(s, ic, 0.5, y, 0.42)
        text(s, 1.05, y, 5.2, 0.42, f"Diff {i+1}", [(t, 13, False, TEXT)], anchor=MSO_ANCHOR.MIDDLE)
    b, tf = box(s, 0.5, 6.42, 5.75, 0.8, "Closes gap", fill=GREEN_T, radius=0.14, pad=0.16, anchor=MSO_ANCHOR.MIDDLE)
    if N8N:
        add_par(tf, [("An OTP can't stop a coached victim. A pause and a family member who can cancel (", 12, False, TEXT),
                     ("n8n", 12, True, N8N_C), (" → Telegram", 12, True, NAVY), (") can, ", 12, False, TEXT),
                     ("used only when the context is risky.", 12, True, GREEN)], first=True)
    else:
      add_par(tf, [("An OTP can't stop a coached victim. A pause and a family member who can cancel can, ", 12, False, TEXT),
                 ("used only when the context is risky.", 12, True, GREEN)], first=True)

    text(s, 6.55, 2.8, 6.3, 0.3, "Scen label", [("THE SAME ₹85,000 PAYMENT · FOUR SITUATIONS", 11, True, NAVY)])
    rows = [(["user", "phone", "keyboard"], "Usual payee, own phone, normal typing", 0.04, "ALLOW", "checkc", GREEN, GREEN_T),
            (["phone", "sim", "pin"], "New phone, fresh SIM, new city", 0.58, "BIOMETRIC", "fp", AMBER, AMBER_T),
            (["phonecall", "clip", "phone"], "On a call, new UPI ID pasted", 0.71, "HOLD + SON", "pause", AMBER, AMBER_T),
            (["bot", "globe", "activity"], "Emulator + VPN, 15 tries in 10 min", 0.97, "BLOCK", "xc", RED, RED_T)]
    for i, (sig, ctx, score, dec, dic, col, tint) in enumerate(rows):
        y = 3.12 + i * 0.8
        b, _ = box(s, 6.55, y, 6.3, 0.72, f"Scenario {i+1}", fill=tint, radius=0.12)
        for k, ic in enumerate(sig):
            icon(s, ic, 6.68 + k * 0.33, y + 0.22, 0.28, col)
        text(s, 7.75, y + 0.06, 2.65, 0.36, f"Scenario {i+1} text", [(ctx, 11.5, True, TEXT)])
        bar(s, 7.75, y + 0.5, 1.8, 0.09, "D9DDE8", f"Risk track {i+1}")
        bar(s, 7.75, y + 0.5, max(0.04, 1.8 * score), 0.09, col, f"Risk fill {i+1}")
        text(s, 9.62, y + 0.4, 0.8, 0.28, f"Risk {i+1}", [(f"risk {score:.2f}", 10, True, col)])
        p = pill(s, 10.45, y + 0.15, 2.28, 0.42, f"Decision {i+1}", dec, col, size=12)
        icon(s, dic, 10.58, y + 0.23, 0.26, WHITE)
    # analyst console: why-bars for situation 3
    b, _ = box(s, 6.55, 6.35, 6.3, 0.9, "Console", fill=NAVY, radius=0.1)
    badge(s, "bars", 6.68, 6.45, 0.4, fill=WHITE, color=NAVY)
    text(s, 7.15, 6.42, 1.55, 0.75, "Console text", [("Why this decision", 11.5, True, WHITE)], anchor=MSO_ANCHOR.TOP)
    text(s, 7.15, 6.85, 1.6, 0.35, "Console sub", [("analyst feedback retrains", 9, False, ICE, True)])
    why = [("On a phone call", 0.31, AMBER), ("New UPI ID pasted", 0.24, AMBER), ("6× usual amount", 0.13, AMBER), ("Own device, home", -0.07, GREEN)]
    for k, (lab, v, col) in enumerate(why):
        y = 6.42 + k * 0.195
        text(s, 8.8, y - 0.02, 1.6, 0.2, f"Why {k+1}", [(lab, 9, False, WHITE)])
        bar(s, 10.4, y + 0.03, 1.75 * abs(v) / 0.31, 0.12, "F0A23B" if col == AMBER else "5BD07A", f"Why bar {k+1}")
        text(s, 12.2, y - 0.02, 0.6, 0.2, f"Why val {k+1}", [(f"{v:+.2f}".replace("-", "−"), 9, True, WHITE)])

# ---------------------------------------------------------------- slide 4
def slide4(s):
    heading(s, headline="> One API call: context in, an explained decision out, in milliseconds")
    top, H = 2.8, 1.95
    boxes = [(0.5, 2.0, "phone", "1 · INPUT", "Payment app + SDK", "amount · device · typing · paste · on-call*"),
             (2.83, 1.8, "server", "2 · API", "/v1/authenticate", "FastAPI · over 150 ms → falls back to 2FA"),
             (4.96, 2.0, "layers", "3 · FEATURES", "20+ features", "vs the user's own baseline: “8× usual amount”")]
    for i, (x, w, ic, tag, title, body) in enumerate(boxes):
        b, _ = box(s, x, top, w, H, f"D{i+1}", fill=PANEL, line=ICE, radius=0.08)
        badge(s, ic, x + 0.12, top + 0.12, 0.45)
        text(s, x + 0.64, top + 0.18, w - 0.7, 0.3, f"D{i+1} tag", [(tag, 10.5, True, MUTED)])
        text(s, x + 0.12, top + 0.66, w - 0.24, 0.5, f"D{i+1} title", [(title, 12.5, True, NAVY)])
        text(s, x + 0.12, top + 1.05, w - 0.24, 0.85, f"D{i+1} body", [(body, 11, False, TEXT)])
    ex, ew = 7.29, 3.15
    b, _ = box(s, ex, top, ew, H, "D4 Engine", fill=NAVY, radius=0.06)
    badge(s, "cpu", ex + 0.12, top + 0.1, 0.36, fill=WHITE, color=NAVY)
    text(s, ex + 0.55, top + 0.13, 2.5, 0.3, "D4 tag", [("4 · RISK ENGINE", 10.5, True, ICE)])
    steps = ["Rules: hard policy + fast-path allow", "LightGBM → fraud probability p", "TreeSHAP → each signal's weight",
             "Risk type: impostor · scam · bot", "Cost policy → decision + challenge"]
    for k, t in enumerate(steps):
        sb, stf = box(s, ex + 0.12, top + 0.55 + k * 0.275, ew - 0.24, 0.24, f"Engine step {k+1}", fill="2C3A86", radius=0.25, pad=0.08, anchor=MSO_ANCHOR.MIDDLE)
        add_par(stf, [("abcde"[k] + "  ", 10, True, ICE), (t, 10, False, WHITE)], first=True, align=L)
    b, _ = box(s, 10.77, top, 2.08, H, "D5", fill=PANEL, line=ICE, radius=0.08)
    text(s, 10.9, top + 0.14, 1.9, 0.3, "D5 tag", [("5 · OUTPUT", 10.5, True, MUTED)])
    for k, (t, col, ic) in enumerate((("ALLOW", GREEN, "checkc"), ("CHALLENGE", AMBER, "fp"), ("BLOCK", RED, "xc"))):
        y = top + 0.5 + k * 0.38
        icon(s, ic, 10.9, y + 0.03, 0.26, col)
        text(s, 11.22, y, 1.6, 0.32, f"D5 out {k+1}", [(t, 12, True, col)], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 10.9, top + 1.62, 1.9, 0.3, "D5 note", [("+ score · reasons · latency", 9.5, False, TEXT)])
    mid = top + 0.9
    arrow(s, 2.5, mid, 2.83, mid, "A1"); arrow(s, 4.63, mid, 4.96, mid, "A2")
    arrow(s, 6.96, mid, 7.29, mid, "A3"); arrow(s, 10.44, mid, 10.77, mid, "A4")

    if N8N:
        n8n_lane(s)
    else:
        b, tf = box(s, 0.5, 4.88, 12.35, 0.4, "Built strip", fill=GREEN_T, radius=0.3, pad=0.14, anchor=MSO_ANCHOR.MIDDLE)
        add_par(tf, [("Built in the hack: ", 11, True, GREEN), ("SDK + demo app · API · risk engine · passkeys · n8n + Telegram loop · dashboard     ", 11, False, TEXT),
                     ("Mocked: ", 11, True, AMBER), ("bank/NPCI rails, *call state", 11, False, TEXT)], first=True, align=L)

    # stack chips
    by = 5.5 if N8N else 5.42
    text(s, 0.5, by, 4.0, 0.3, "Stack label", [("STACK · ₹0 RUNNING COST", 11, True, NAVY)])
    chips = [("code", "Python · FastAPI"), ("cpu", "LightGBM + SHAP"), ("eye", "React dashboard"),
             ("db", "Supabase"), ("send", "n8n + Telegram"), ("fp", "WebAuthn passkeys")]
    for k, (ic, t) in enumerate(chips):
        x = 0.5 + (k % 2) * 2.03; y = by + 0.36 + (k // 2) * 0.5
        hi = N8N and t.startswith("n8n")
        b, _ = box(s, x, y, 1.95, 0.42, f"Chip {k+1}", fill=N8N_T if hi else PANEL, line=N8N_C if hi else None, radius=0.3)
        icon(s, ic, x + 0.12, y + 0.08, 0.26, N8N_C if hi else NAVY)
        runs = [("n8n", 10.5, True, N8N_C), (t[3:], 10.5, True, TEXT)] if hi else [(t, 10.5, True, TEXT)]
        text(s, x + 0.44, y, 1.5, 0.42, f"Chip {k+1} text", runs, anchor=MSO_ANCHOR.MIDDLE)
    # data tiles
    text(s, 4.68, by, 3.9, 0.3, "Data label", [("SYNTHETIC DATA (NO PUBLIC SET HAS THESE SIGNALS)", 11, True, NAVY)])
    for k, (big, small) in enumerate((("5k", "users"), ("30", "days"), ("4", "attack types"))):
        x = 4.68 + k * 1.31
        b, _ = box(s, x, by + 0.36, 1.2, 0.92, f"Data {k+1}", fill=PANEL, radius=0.12)
        text(s, x, by + 0.4, 1.2, 0.5, f"Data {k+1} big", [(big, 24, True, NAVY, False, HEAD)], align=C)
        text(s, x, by + 0.9, 1.2, 0.3, f"Data {k+1} small", [(small, 10.5, False, TEXT)], align=C)
    text(s, 4.68, by + 1.36, 3.9, 0.4, "Data note", [("+ hard negatives (travel, new phone) · checked on IEEE-CIS & PaySim", 10, False, MUTED, True)])
    # decision rule chart
    amts = ["₹500", "₹1k", "₹2k", "₹5k", "₹10k", "₹50k", "₹85k"]
    vals = [min(0.6, max(0.05, 500 / a)) for a in (500, 1000, 2000, 5000, 10000, 50000, 85000)]
    cd = CategoryChartData(); cd.categories = amts
    cd.add_series("challenge threshold", vals); cd.add_series("block threshold", [0.9] * len(amts))
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(8.8), Inches(by - 0.05), Inches(4.05), Inches(1.88), cd)
    gf.name = "Decision rule chart"
    ch = gf.chart
    ch.has_title = True; ch.chart_title.text_frame.text = "Decision rule: stricter as the amount grows"
    tp = ch.chart_title.text_frame.paragraphs[0]; tp.runs[0].font.size = Pt(11); tp.runs[0].font.bold = True; tp.runs[0].font.color.rgb = rgb(NAVY); tp.runs[0].font.name = BODY
    ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(9); ch.legend.font.color.rgb = rgb(MUTED)
    va = ch.value_axis; va.maximum_scale = 1.0; va.minimum_scale = 0; va.major_unit = 0.5
    va.tick_labels.font.size = Pt(9); va.tick_labels.font.color.rgb = rgb(MUTED)
    va.major_gridlines.format.line.color.rgb = rgb("E3E6EF"); va.format.line.fill.background()
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.color.rgb = rgb(MUTED)
    ca.format.line.color.rgb = rgb("C9CEDB")
    for ser, col in zip(ch.plots[0].series, (AMBER, RED)):
        ser.format.line.color.rgb = rgb(col); ser.format.line.width = Pt(2.25); ser.smooth = False
        ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb = rgb(col); ser.marker.format.line.fill.background()
        ser.marker.size = 5

def n8n_lane(s):
    """Where n8n is used: the guardian loop that runs after a HOLD decision, off the scoring path."""
    y, h = 4.83, 0.48
    b, tf = box(s, 0.5, y, 2.05, h, "n8n head", fill=N8N_C, radius=0.2, pad=0.08, anchor=MSO_ANCHOR.MIDDLE)
    tf.margin_left = Inches(0.46)
    add_par(tf, [("POWERED BY n8n", 11, True, WHITE)], first=True, align=L)
    add_par(tf, [("guardian loop", 8.5, False, WHITE, True)], align=L)
    icon(s, "refresh", 0.6, y + 0.13, 0.26, WHITE)
    steps = [("pause", AMBER, "Scam-risk HOLD"), ("zap", NAVY, "n8n workflow triggered"),
             ("send", NAVY, "Telegram: Approve / Cancel"), ("clock", NAVY, "No reply in 30 min → cancel"),
             ("db", GREEN, "Outcome → training label")]
    x, w, gap = 2.75, 1.884, 0.17
    for k, (ic, col, t) in enumerate(steps):
        hi = t.startswith("n8n")
        b, _ = box(s, x, y, w, h, f"n8n step {k+1}", fill=N8N_T if hi else PANEL, line=N8N_C if hi else ICE, radius=0.2)
        icon(s, ic, x + 0.1, y + 0.12, 0.26, N8N_C if hi else col)
        runs = [("n8n", 10, True, N8N_C), (t[3:], 10, True, TEXT)] if hi else [(t, 10, True, TEXT)]
        text(s, x + 0.42, y, w - 0.46, h, f"n8n step {k+1} text", runs, anchor=MSO_ANCHOR.MIDDLE)
        if k < 4:
            arrow(s, x + w, y + h / 2, x + w + gap, y + h / 2, f"n8n arrow {k+1}", width=1.25)
        x += w + gap

# ---------------------------------------------------------------- slide 5
def slide5(s):
    heading(s, headline="> A deployed, working prototype in the hack, honest about its limits")
    cols = [("1 · WHAT WE BUILD", 0.5, "code"), ("2 · RISKS & HOW WE HANDLE THEM", 4.69, "alert"), ("3 · EXPECTED IMPACT", 8.88, "target")]
    for i, (h, x, ic) in enumerate(cols):
        b, tf = box(s, x, 2.8, 3.97, 0.46, f"Col head {i+1}", fill=NAVY, radius=0.3, pad=0.12, anchor=MSO_ANCHOR.MIDDLE)
        tf.margin_left = Inches(0.5)
        add_par(tf, [(h, 12, True, WHITE)], first=True, align=L)
        icon(s, ic, x + 0.15, 2.9, 0.26, WHITE)
    # col 1: checklist + timeline
    b, _ = box(s, 0.5, 3.36, 3.97, 3.86, "Build body", fill=PANEL, radius=0.05)
    items = ["Live API with fail-safe to 2FA", "Rules + LightGBM + TreeSHAP", "Passkey fingerprint check", "n8n + Telegram trusted-contact loop" if N8N else "Telegram trusted-contact loop", "Dashboard with “why” panel"]
    for k, t in enumerate(items):
        y = 3.5 + k * 0.44
        badge(s, "check", 0.65, y + 0.04, 0.3, fill=GREEN, ratio=0.6)
        runs = [("n8n", 12.5, True, N8N_C), (t[3:], 12.5, False, TEXT)] if N8N and t.startswith("n8n") else [(t, 12.5, False, TEXT)]
        text(s, 1.07, y, 3.3, 0.38, f"Build {k+1}", runs, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.65, 5.78, 3.7, 0.3, "Timeline label", [("24-HOUR PLAN", 10.5, True, NAVY)])
    segs = [("Data", "H0–6", "3F51B5"), ("Engine", "H6–12", "2C3A86"), ("App + UI", "H12–18", "1E2761"), ("Deploy", "H18–24", GREEN)]
    for k, (t, h, col) in enumerate(segs):
        x = 0.65 + k * 0.92
        b, tf = box(s, x, 6.08, 0.9, 0.42, f"Seg {k+1}", fill=col, radius=0.15, pad=0.03, anchor=MSO_ANCHOR.MIDDLE)
        add_par(tf, [(t, 10, True, WHITE)], first=True, align=C)
        text(s, x, 6.52, 0.9, 0.25, f"Seg {k+1} h", [(h, 9, False, MUTED)], align=C)
    text(s, 0.65, 6.82, 3.7, 0.3, "Cost", [("Cost: ", 11, True, NAVY), ("₹0, free tiers + open source", 11, False, TEXT)])
    # col 2: risks with icons
    risks = [("db", "Synthetic data", "hard negatives, per-type metrics, public-set checks"),
             ("radio", "Call state needs an SDK", "missing = “unknown”, never “safe”"),
             ("alert", "False positives", "block only at p ≥ 0.9; ≤ 5% challenged"),
             ("lock", "Privacy (DPDP)", "derived features only, hashed IDs")]
    for k, (ic, r, m) in enumerate(risks):
        y = 3.36 + k * 0.98
        b, _ = box(s, 4.69, y, 3.97, 0.88, f"Risk {k+1}", fill=PANEL, radius=0.1)
        badge(s, ic, 4.82, y + 0.2, 0.48, fill=RED_T, color=RED)
        text(s, 5.42, y + 0.1, 3.15, 0.3, f"Risk {k+1} head", [(r, 12.5, True, RED)])
        text(s, 5.42, y + 0.42, 3.15, 0.42, f"Risk {k+1} fix", [("→ " + m, 11, False, TEXT)])
    # col 3: tagline + metrics
    b, tf = box(s, 8.88, 3.36, 3.97, 0.62, "Impact lead", fill=GREEN_T, radius=0.12, pad=0.15, anchor=MSO_ANCHOR.MIDDLE)
    add_par(tf, [("Honest payments stay instant.", 12, True, GREEN)], first=True, align=C)
    add_par(tf, [("Risky ones get stopped.", 12, True, GREEN)], align=C)
    mets = [("shield", "≥ 80%", "scams caught, ≤ 5% of honest payments challenged"),
            ("zap", "≥ 90%", "honest ₹10k+ transfers with no delay"),
            ("gauge", "< 50 ms", "p95 decision time (UPI allows 15 s)")]
    for k, (ic, big, small) in enumerate(mets):
        y = 4.08 + k * 0.92
        b, _ = box(s, 8.88, y, 3.97, 0.82, f"Metric {k+1}", fill=WHITE, line=ICE, radius=0.1)
        badge(s, ic, 9.0, y + 0.17, 0.48, fill=GREEN)
        text(s, 9.58, y + 0.04, 3.2, 0.42, f"Metric {k+1} value", [(big, 20, True, GREEN, False, HEAD)])
        text(s, 9.58, y + 0.46, 3.2, 0.34, f"Metric {k+1} text", [(small, 10.5, False, TEXT)])
    text(s, 8.88, 6.88, 3.97, 0.3, "Tagline", [("Fast. Secure. Explained.", 12, True, NAVY, True, HEAD)], align=C)

def main(src, out):
    prs = Presentation(src)
    sl = list(prs.slides)
    for s in sl:
        for sh in list(s.shapes):
            if not sh.name.startswith(("Text ", "Shape ")):
                remove(sh)
    slide1(sl[0]); slide2(sl[1]); slide3(sl[2]); slide4(sl[3]); slide5(sl[4])
    prs.save(out)

if __name__ == "__main__":
    ICON_DIR = sys.argv[3]
    N8N = "--n8n" in sys.argv[4:]
    main(sys.argv[1], sys.argv[2])
